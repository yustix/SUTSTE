"""Genera documentos de muestra para probar SUT STE."""
import os
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas
from docx import Document
from docx.shared import Pt

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "muestras")
OUT = os.path.abspath(OUT)
os.makedirs(OUT, exist_ok=True)


def pdf_oficio():
    path = os.path.join(OUT, "oficio-solicitud-plazo.pdf")
    c = canvas.Canvas(path, pagesize=LETTER)
    w, h = LETTER
    c.setFont("Helvetica-Bold", 12)
    c.drawString(2.5 * cm, h - 2.5 * cm, "SINDICATO ÚNICO DE TRABAJADORES")
    c.setFont("Helvetica", 10)
    c.drawString(2.5 * cm, h - 3.1 * cm, "Secretaría General - Ciudad de México")
    c.line(2.5 * cm, h - 3.4 * cm, w - 2.5 * cm, h - 3.4 * cm)

    c.setFont("Helvetica-Bold", 10)
    c.drawString(2.5 * cm, h - 4.4 * cm, "OFICIO NUM. SG/147/2026")
    c.drawString(2.5 * cm, h - 5.0 * cm, "ASUNTO:")
    c.setFont("Helvetica", 10)
    c.drawString(4.6 * cm, h - 5.0 * cm, "Solicitud de información sobre tabulador salarial")
    c.setFont("Helvetica-Bold", 10)
    c.drawString(2.5 * cm, h - 5.6 * cm, "FECHA:")
    c.setFont("Helvetica", 10)
    c.drawString(4.6 * cm, h - 5.6 * cm, "Ciudad de Mexico, a 8 de septiembre de 2026")

    c.setFont("Helvetica-Bold", 10)
    y = h - 7.0 * cm
    c.drawString(2.5 * cm, y, "C. LIC. MARIA ELENA ROBLES CASTILLO")
    c.setFont("Helvetica", 10)
    c.drawString(2.5 * cm, y - 0.5 * cm, "Directora de Recursos Humanos")
    c.drawString(2.5 * cm, y - 1.0 * cm, "PRESENTE.")

    cuerpo = [
        "Por medio del presente, y en atencion a los acuerdos de la mesa de trabajo celebrada el",
        "pasado 25 de agosto, solicito a usted tenga a bien remitir la informacion correspondiente",
        "al tabulador salarial vigente del personal operativo, asi como el detalle del incremento",
        "autorizado para el ejercicio 2026.",
        "",
        "Se requiere que la documentacion se envie a mas tardar el 22 de septiembre de 2026, con el",
        "objeto de integrarla al expediente de la revision contractual. El monto observado en el",
        "ultimo informe asciende a $1,250,000.00 (un millon doscientos cincuenta mil pesos 00/100 M.N.),",
        "cantidad que debera conciliarse con la tesoreria.",
        "",
        "Asimismo, se le solicita confirmar la fecha de la proxima reunion de la comision mixta,",
        "a fin de convocar a los representantes de las areas involucradas. En caso de existir algun",
        "inconveniente para cumplir con el plazo senalado, agradecemos nos lo comunique por escrito",
        "dentro de los tres dias habiles siguientes a la recepcion del presente.",
        "",
        "Sin otro particular por el momento, y agradeciendo de antemano la atencion brindada, quedo",
        "de usted para cualquier aclaracion.",
    ]
    c.setFont("Helvetica", 10)
    y -= 2.0 * cm
    for linea in cuerpo:
        c.drawString(2.5 * cm, y, linea)
        y -= 0.52 * cm

    y -= 1.0 * cm
    c.setFont("Helvetica-Bold", 10)
    c.drawString(2.5 * cm, y, "ATENTAMENTE")
    y -= 1.6 * cm
    c.drawString(2.5 * cm, y, "____________________________")
    y -= 0.55 * cm
    c.setFont("Helvetica", 10)
    c.drawString(2.5 * cm, y, "ING. JORGE LUIS MENDOZA ARRIAGA")
    y -= 0.5 * cm
    c.drawString(2.5 * cm, y, "Secretario General")
    y -= 0.9 * cm
    c.setFont("Helvetica", 8)
    c.drawString(2.5 * cm, y, "c.c.p. Archivo. Tesoreria. Contacto: oficinapartes@sutste.example.mx - Tel. 55 5555 1234")
    c.save()
    print("generado", path)


def docx_minuta():
    path = os.path.join(OUT, "minuta-mesa-de-trabajo.docx")
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)

    d.add_paragraph("MINUTA DE LA TERCERA MESA DE TRABAJO")
    d.add_paragraph("Sindicato Unico de Trabajadores - Direccion de Recursos Humanos")
    d.add_paragraph("Ciudad de Mexico, siendo las 10:00 horas del 3 de septiembre de 2026.")
    d.add_paragraph("ASUNTO: Seguimiento a la revision contractual y al pago de prestaciones pendientes.")
    d.add_paragraph("ASISTENTES: Ing. Jorge Luis Mendoza Arriaga, Secretario General; Lic. Maria Elena Robles Castillo, Directora de Recursos Humanos; C. Ana Luisa Fuentes Ortega, Secretaria de Finanzas.")
    d.add_paragraph("")
    d.add_paragraph("ORDEN DEL DIA")
    for t in [
        "1. Lista de asistencia y declaracion de quorum.",
        "2. Lectura y aprobacion de la minuta anterior.",
        "3. Estado que guarda el pago del bono de despensa correspondiente al segundo semestre.",
        "4. Revision del tabulador salarial y del incremento autorizado para 2026.",
        "5. Asuntos generales.",
    ]:
        d.add_paragraph(t)
    d.add_paragraph("")
    d.add_paragraph("ACUERDOS")
    for t in [
        "Primero. La Direccion de Recursos Humanos entregara el detalle del bono de despensa a mas tardar el 30 de septiembre de 2026, por un monto estimado de $480,000.00 (cuatrocientos ochenta mil pesos 00/100 M.N.).",
        "Segundo. El sindicato remitira por escrito las observaciones al tabulador salarial dentro de un plazo de cinco dias habiles contados a partir de la recepcion del documento.",
        "Tercero. Se convoca a la cuarta mesa de trabajo para el 15 de octubre de 2026 a las 10:00 horas, en la sala de juntas de la Direccion.",
        "Cuarto. Ambas partes se comprometen a dar respuesta a las solicitudes pendientes antes del cierre del ejercicio fiscal.",
    ]:
        d.add_paragraph(t)
    d.add_paragraph("")
    d.add_paragraph("No habiendo mas asuntos que tratar, se da por concluida la reunion siendo las 12:30 horas, firmando al margen y al calce los que en ella intervinieron.")
    d.add_paragraph("")
    d.add_paragraph("ATENTAMENTE")
    d.add_paragraph("ING. JORGE LUIS MENDOZA ARRIAGA - Secretario General")
    d.save(path)
    print("generado", path)


def txt_circular():
    path = os.path.join(OUT, "circular-informativa.txt")
    texto = """CIRCULAR NUM. 12/2026

ASUNTO: Aviso sobre el horario de atencion en ventanilla durante el periodo vacacional.

A TODO EL PERSONAL AGREMIADO:

Se hace de su conocimiento que, con motivo del periodo vacacional de fin de ano, la atencion en ventanilla se brindara de las 9:00 a las 14:00 horas, del 22 de diciembre de 2026 al 2 de enero de 2027.

La presente circular tiene caracter informativo, por lo que no requiere respuesta alguna. Los tramites iniciados antes del 19 de diciembre continuaran su curso habitual a partir del 5 de enero de 2027.

Para dudas sobre el calendario, se encuentra disponible el aviso en el tablero de la sede y en el correo institucional.

ATENTAMENTE
La Secretaria de Actas y Acuerdos
"""
    with open(path, "w", encoding="utf-8") as f:
        f.write(texto)
    print("generado", path)


pdf_oficio()
docx_minuta()
txt_circular()
print("listo en", OUT)
